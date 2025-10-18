#!/usr/bin/env python3
"""
analytics_plots.py
- Runs diagnostic queries against your postgres DB
- Runs analytics queries, saves CSVs and an Excel workbook
- Creates 6 plots (handles empty data by saving placeholder images)
"""

from pathlib import Path
import pandas as pd
import matplotlib.pyplot as plt
from sqlalchemy import create_engine
import sys
import plotly.express as px
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# --- DB settings ---
DB = {
    'dbname': 'steamdb',
    'user': 'steam_user',
    'password': 'steam_pass',
    'host': 'localhost',
    'port': 5432
}

CHARTS_DIR = Path("charts")
CHARTS_DIR.mkdir(parents=True, exist_ok=True)

EXPORTS_DIR = Path("exports")
EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_DIR = Path("data/outputs")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# --- queries ---
QUERIES = {
    "revenue_by_year": """
        SELECT DATE_TRUNC('year', purchase_date) AS year, SUM(price) AS total_revenue
        FROM purchases
        GROUP BY year
        ORDER BY year;
    """,
    "avg_playtime_by_genre": """
        SELECT ge.name AS genre, AVG(g.average_playtime) AS avg_playtime
        FROM game_genres gg
        JOIN genres ge ON gg.genre_id = ge.genre_id
        JOIN games g ON gg.game_id = g.game_id
        GROUP BY ge.name
        ORDER BY avg_playtime DESC
        LIMIT 10;
    """,
    "purchases_by_country": """
        SELECT u.country, COUNT(p.purchase_id) AS purchases
        FROM purchases p
        JOIN users u ON p.user_id = u.user_id
        GROUP BY u.country
        ORDER BY purchases DESC
        LIMIT 15;
    """,
    "top10_games_by_purchases": """
        SELECT g.name, COUNT(p.purchase_id) AS purchases_count
        FROM purchases p
        JOIN games g ON p.game_id = g.game_id
        GROUP BY g.name
        ORDER BY purchases_count DESC
        LIMIT 10;
    """,
    "top_developers_by_revenue": """
        SELECT d.name AS developer, COALESCE(SUM(p.price),0) AS revenue
        FROM purchases p
        JOIN games g ON p.game_id = g.game_id
        JOIN developers d ON g.developer_id = d.developer_id
        GROUP BY d.name
        ORDER BY revenue DESC
        LIMIT 10;
    """,
    "daily_purchases_30days": """
        SELECT DATE(purchase_date) AS day, COUNT(*) AS purchases
        FROM purchases
        WHERE purchase_date >= (CURRENT_DATE - INTERVAL '30 days')
        GROUP BY day
        ORDER BY day;
    """
}

DIAGNOSTIC_QUERIES = {
    "count_games": "SELECT COUNT(*) FROM games;",
    "count_purchases": "SELECT COUNT(*) FROM purchases;",
    "min_max_purchase_date": "SELECT MIN(purchase_date) AS min_date, MAX(purchase_date) AS max_date FROM purchases;",
    "price_stats": "SELECT COUNT(price) AS cnt, SUM(price) AS sum_price, AVG(price) AS avg_price FROM purchases;"
}

# --- Helpers ---
def get_engine():
    url = f"postgresql+psycopg2://{DB['user']}:{DB['password']}@{DB['host']}:{DB['port']}/{DB['dbname']}"
    return create_engine(url, echo=False)

def run_sql_to_df(sql, engine):
    try:
        return pd.read_sql(sql, engine)
    except Exception as ex:
        print("SQL read error:", ex)
        return pd.DataFrame()

def save_csv(df, name):
    file = OUTPUT_DIR / f"{name}.csv"
    df.to_csv(file, index=False, encoding="utf-8")
    print(f"[OK] Saved CSV: {file}")
    return file

# === UPDATED EXCEL EXPORT ===
def save_excel(all_dfs, filename):
    file = EXPORTS_DIR / filename
    with pd.ExcelWriter(file, engine="openpyxl") as writer:
        for name, df in all_dfs.items():
            sheet = name[:31]  # Excel sheet name limit
            try:
                df.to_excel(writer, sheet_name=sheet, index=False)
            except Exception as ex:
                print(f"[WARN] Could not write sheet {sheet}: {ex}")

    wb = load_workbook(file)
    for sheet in wb.sheetnames:
        ws = wb[sheet]

        ws.freeze_panes = "A2"

        max_col = ws.max_column
        max_row = ws.max_row
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

        for col in range(1, max_col + 1):
            col_letter = get_column_letter(col)
            values = [cell.value for cell in ws[col_letter][1:]] 
            if all(isinstance(v, (int, float)) and v is not None for v in values):
                color_scale = ColorScaleRule(
                    start_type="min", start_color="FFFFFF",
                    mid_type="percentile", mid_value=50, mid_color="FFFF99",
                    end_type="max", end_color="FF0000"
                )
                ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{max_row}", color_scale)

    wb.save(file)
    print(f"[OK] Excel report saved: {file}")
    return file

# --- Plotting helpers ---
def create_placeholder_image(filepath, title="No data", text="No data to plot"):
    fig = plt.figure(figsize=(8,4))
    plt.text(0.5, 0.5, text, ha='center', va='center', fontsize=14)
    plt.title(title)
    plt.axis('off')
    fig.tight_layout()
    fig.savefig(filepath)
    plt.close(fig)
    print(f"[OK] Placeholder image saved: {filepath}")

def save_line_plot(df, x_col, y_col, title, fname, xlabel=None, ylabel=None):
    file = CHARTS_DIR / fname
    if df.empty or df[y_col].isna().all() or (df[y_col].sum() == 0):
        create_placeholder_image(file, title, "No data to plot")
        return file
    fig, ax = plt.subplots(figsize=(10,5))
    ax.plot(df[x_col], df[y_col], marker="o")
    ax.set_title(title)
    ax.set_xlabel(xlabel or x_col)
    ax.set_ylabel(ylabel or y_col)
    ax.grid(True)
    fig.tight_layout()
    fig.savefig(file)
    plt.close(fig)
    print(f"[OK] Saved: {file}")
    return file

def save_barh_plot(df, label_col, value_col, title, fname, top_n=10):
    file = CHARTS_DIR / fname
    if df.empty or df[value_col].isna().all() or (df[value_col].sum() == 0):
        create_placeholder_image(file, title, "No data to plot")
        return file
    df_plot = df.sort_values(value_col, ascending=True).tail(top_n)
    fig, ax = plt.subplots(figsize=(10,5))
    ax.barh(df_plot[label_col].astype(str), df_plot[value_col])
    ax.set_title(title)
    ax.set_xlabel(value_col)
    fig.tight_layout()
    fig.savefig(file)
    plt.close(fig)
    print(f"[OK] Saved: {file}")
    return file

def save_pie_plot(df, label_col, value_col, title, fname, top_n=10):
    file = CHARTS_DIR / fname
    if df.empty or df[value_col].isna().all() or (df[value_col].sum() == 0):
        create_placeholder_image(file, title, "No data to plot")
        return file
    df_plot = df.head(top_n)
    fig, ax = plt.subplots(figsize=(7,7))
    ax.pie(df_plot[value_col], labels=df_plot[label_col].astype(str), autopct="%1.1f%%")
    ax.set_title(title)
    fig.tight_layout()
    fig.savefig(file)
    plt.close(fig)
    print(f"[OK] Saved: {file}")
    return file

# --- Main ---
def main():
    engine = get_engine()

    # Diagnostics
    print("=== DIAGNOSTICS ===")
    for name, q in DIAGNOSTIC_QUERIES.items():
        df = run_sql_to_df(q, engine)
        print(f"{name}:")
        print(df)

    # Run analytics queries
    all_results = {}
    for name, q in QUERIES.items():
        df = run_sql_to_df(q, engine)
        save_csv(df, name)
        all_results[name] = df

    # Save Excel
    save_excel(all_results, "report.xlsx")

    # Plots
    save_line_plot(run_sql_to_df(QUERIES["revenue_by_year"], engine), "year", "total_revenue",
                   "Revenue by Year", "revenue_by_year.png", xlabel="Year", ylabel="Total Revenue")

    save_barh_plot(run_sql_to_df(QUERIES["avg_playtime_by_genre"], engine), "genre", "avg_playtime",
                   "Average Playtime by Genre", "avg_playtime_by_genre.png")

    save_pie_plot(run_sql_to_df(QUERIES["purchases_by_country"], engine), "country", "purchases",
                  "Purchases by Country", "purchases_by_country.png")

    save_barh_plot(run_sql_to_df(QUERIES["top10_games_by_purchases"], engine), "name", "purchases_count",
                   "Top 10 Games by Purchases", "top10_games.png")

    save_barh_plot(run_sql_to_df(QUERIES["top_developers_by_revenue"], engine), "developer", "revenue",
                   "Top Developers by Revenue", "top_developers_by_revenue.png")

    save_line_plot(run_sql_to_df(QUERIES["daily_purchases_30days"], engine), "day", "purchases",
                   "Daily Purchases (Last 30 Days)", "daily_purchases_30days.png", xlabel="Day", ylabel="Purchases")

    print("=== DONE ===")

if __name__ == "__main__":
    main()
